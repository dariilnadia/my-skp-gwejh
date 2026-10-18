import streamlit as st
import pandas as pd
from datetime import datetime
import os
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- KONFIGURASI HALAMAN ---
st.set_page_config(
    page_title="SKP & Eviden - Serang Gawe FM",
    page_icon="🎙️",
    layout="wide"
)

# --- TAMPILAN KUSTOM (TEMA BIRU SERANG GAWE FM) ---
st.markdown("""
    <style>
    .stApp {
        background-color: #f4f7fc;
    }
    .main-header {
        background-color: #0d47a1;
        color: white;
        padding: 20px;
        border-radius: 10px;
        text-align: center;
        margin-bottom: 25px;
        box-shadow: 0 4px 6px rgba(0,0,0,0.1);
    }
    .stButton>button {
        background-color: #1565c0;
        color: white;
        font-weight: bold;
        border-radius: 8px;
        padding: 10px 24px;
        border: none;
        width: 100%;
    }
    .stButton>button:hover {
        background-color: #0d47a1;
        color: white;
    }
    </style>
""", unsafe_allow_html=True)

# --- HEADER APLIKASI ---
st.markdown("""
    <div class="main-header">
        <h1>🎙️ PORTAL EVIDEN SKP - SERANG GAWE FM</h1>
        <p>Sistem Rekapitulasi Dokumentasi & Laporan Kinerja Bulanan Karyawan</p>
    </div>
""", unsafe_allow_html=True)

# Folder penyimpanan foto
if not os.path.exists("uploads"):
    os.makedirs("uploads")

# File database lokal
DATA_FILE = "data_eviden.csv"

if os.path.exists(DATA_FILE):
    df_data = pd.read_csv(DATA_FILE)
else:
    df_data = pd.DataFrame(columns=["Tanggal", "Waktu", "Program", "Narasumber", "Tema", "Foto_Path"])

# --- HELPER FUNCTION: BORDER TABEL ---
def set_cell_border(cell, **kwargs):
    """Memberikan border/garis tepi hitam pada sel tabel Word."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, attr in [('val', 'w:val'), ('color', 'w:color'), ('sz', 'w:sz'), ('space', 'w:space')]:
                if key in edge_data:
                    element.set(qn(attr), str(edge_data[key]))

def set_table_borders(table):
    """Menambahkan border grid penuh untuk tabel."""
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'  <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(borders)

# --- MENU TAB ---
tab1, tab2, tab3 = st.tabs(["📝 Input Eviden Harian", "📊 Rekap Data Bulanan", "📄 Generate Laporan SKP"])

# --- TAB 1: FORM INPUT ---
with tab1:
    st.subheader("📌 Form Input Kegiatan Harian")
    
    with st.form("form_eviden", clear_on_submit=True):
        col1, col2 = st.columns(2)
        with col1:
            tanggal = st.date_input("Tanggal Kegiatan", datetime.now())
            waktu = st.selectbox("Waktu/Sesi Siaran", [
                "SIARAN PAGI", 
                "OBRAS", 
                "MENJELANG SORE", 
                "SERGAW MALAM", 
                "NGECAS IMAN", 
                "PROSES PENCARIAN NARASUMBER", 
                "PROSES MEMBUAT RANGKUMAN LIVE"
            ])
            program = st.text_input("Nama Program / Jenis Kegiatan", placeholder="Contoh: Zona Atensi / OBRAS / Siaran Pagi")
        
        with col2:
            narasumber = st.text_area("Nama Narasumber & Jabatan", placeholder="Contoh: Drs. Lutfi Mujahidin, M.Si. (Kepala BPBD)")
            tema = st.text_area("Tema / Judul Pembahasan", placeholder="Contoh: Status Siaga Darurat Bencana")
        
        uploaded_file = st.file_uploader("Upload Foto Dokumentasi/Evidence (PNG/JPG)", type=["jpg", "jpeg", "png"])
        
        submit = st.form_submit_button("💾 Simpan Eviden")
        
        if submit:
            if uploaded_file is not None:
                img_name = f"{tanggal.strftime('%Y%m%d')}_{uploaded_file.name}"
                img_path = os.path.join("uploads", img_name)
                with open(img_path, "wb") as f:
                    f.write(uploaded_file.getbuffer())
                
                new_entry = {
                    "Tanggal": tanggal.strftime("%d %B %Y").upper(),
                    "Waktu": waktu,
                    "Program": program,
                    "Narasumber": narasumber,
                    "Tema": tema,
                    "Foto_Path": img_path
                }
                df_data = pd.concat([df_data, pd.DataFrame([new_entry])], ignore_index=True)
                df_data.to_csv(DATA_FILE, index=False)
                st.success("✅ Data eviden berhasil disimpan!")
            else:
                st.error("⚠️ Harap upload foto bukti dokumentasi!")

# --- TAB 2: REKAP DATA ---
with tab2:
    st.subheader("📋 Rekapitulasi Data Eviden")
    if not df_data.empty:
        st.dataframe(df_data[["Tanggal", "Waktu", "Program", "Narasumber", "Tema"]], use_container_width=True)
    else:
        st.info("Belum ada data yang diinputkan bulan ini.")

# --- TAB 3: GENERATE LAPORAN DOCX ---
with tab3:
    st.subheader("🖨️ Cetak Laporan Bulanan (Format SKP Presisi)")
    
    col_a, col_b = st.columns(2)
    with col_a:
        nama_pegawai = st.text_input("Nama Pegawai", value="NADIYYATUL FIZRIYA")
        nip_pegawai = st.text_input("NIP / ID Pegawai", value="200112222025212026")
    with col_b:
        bulan_tahun = st.text_input("Bulan & Tahun Laporan", value="JANUARI 2026")
    
    if st.button("🚀 Buat Laporan Word (.docx) Otomatis"):
        if not df_data.empty:
            doc = Document()
            
            # --- HALAMAN 1: SAMPUL ---
            p_cover = doc.add_paragraph()
            p_cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            run1 = p_cover.add_run("LAMPIRAN\nSKP BULAN " + bulan_tahun.upper() + "\n\n\n\n\n\n")
            run1.font.size = Pt(16)
            run1.font.bold = True
            
            run_nama = p_cover.add_run(nama_pegawai.upper() + "\n")
            run_nama.font.size = Pt(18)
            run_nama.font.bold = True
            
            run_nip = p_cover.add_run(nip_pegawai)
            run_nip.font.size = Pt(14)
            
            doc.add_page_break()
            
            # --- HALAMAN 2: DOKUMENTASI GRID DENGAN BORDER KOTAK ---
            p_head = doc.add_paragraph()
            p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_h = p_head.add_run(f"DOKUMENTASI {bulan_tahun.upper()}\n")
            run_h.font.bold = True
            run_h.font.size = Pt(14)
            
            p_narrative = doc.add_paragraph(
                "Keberhasilan pelaksanaan kegiatan sosialisasi kepada masyarakat di wilayah Kabupaten Serang tidak terlepas dari alur persiapan dan koordinasi intensif yang dilakukan di balik layar. Seluruh dokumen dan foto yang dilampirkan dalam bagian ini merupakan rekapitulasi dari peran operasional teknis yang telah dilaksanakan secara berkesinambungan.\n\n"
                "Evidence ini mencakup aktivitas riset dan komunikasi dalam menjaring narasumber yang relevan, pengelolaan administratif, hingga dukungan operasional teknis acara. Penulisan dan penyampaian eviden ini disusun untuk memberikan gambaran utuh mengenai kontribusi nyata serta realisasi sasaran kinerja pegawai pada periode ini.\n"
            )
            p_narrative.paragraph_format.line_spacing = 1.15
            p_narrative.paragraph_format.space_after = Pt(14)
            
            # Grid 3 Kolom untuk Frame Foto
            grid_table = doc.add_table(rows=0, cols=3)
            grid_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            
            row_cells = None
            for idx, row in df_data.iterrows():
                if idx % 3 == 0:
                    row_cells = grid_table.add_row().cells
                col_idx = idx % 3
                cell = row_cells[col_idx]
                
                # Buat tabel kecil 1x1 di dalam sel untuk efek BORDER FRAME KOTAK
                inner_table = cell.add_table(rows=1, cols=1)
                inner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
                inner_cell = inner_table.rows[0].cells[0]
                
                # Pasang border hitam di sekeliling frame
                border_spec = {"val": "single", "sz": "8", "color": "000000"}
                set_cell_border(inner_cell, top=border_spec, bottom=border_spec, left=border_spec, right=border_spec)
                
                p_frame = inner_cell.paragraphs[0]
                p_frame.alignment = WD_ALIGN_PARAGRAPH.CENTER
                
                # Teks Judul Foto di Dalam Frame
                if "PROSES" in str(row['Waktu']).upper():
                    label_text = f"{row['Waktu']}"
                else:
                    label_text = f"{row['Tanggal']}\n({row['Waktu']})"
                    
                r_title = p_frame.add_run(f"{label_text}\n\n")
                r_title.font.bold = True
                r_title.font.size = Pt(9.5)
                
                # Masukkan Gambar
                if os.path.exists(row['Foto_Path']):
                    p_frame.add_run().add_picture(row['Foto_Path'], width=Inches(1.6))
                    
                # Kasih spasi bawah di luar inner table
                cell.add_paragraph().paragraph_format.space_after = Pt(8)
            
            doc.add_page_break()
            
            # --- HALAMAN 3: TABEL SKP RAPI ---
            doc.add_heading(f"REKAPITULASI SIARAN & KEGIATAN", level=2)
            
            skp_table = doc.add_table(rows=1, cols=5)
            skp_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(skp_table)
            
            hdr_cells = skp_table.rows[0].cells
            headers = ['NO', 'HARI/TANGGAL', 'NAMA PROGRAM', 'NARASUMBER', 'TEMA']
            for i, h in enumerate(headers):
                hdr_cells[i].text = h
                p = hdr_cells[i].paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.bold = True
                    r.font.size = Pt(10)
            
            for idx, row in df_data.iterrows():
                row_cells = skp_table.add_row().cells
                row_cells[0].text = str(idx + 1)
                row_cells[1].text = str(row['Tanggal'])
                row_cells[2].text = f"{row['Program']}\n({row['Waktu']})" if row['Program'] else str(row['Waktu'])
                row_cells[3].text = str(row['Narasumber']) if pd.notna(row['Narasumber']) else "-"
                row_cells[4].text = str(row['Tema']) if pd.notna(row['Tema']) else "-"
                
                # Formatting sel
                row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
                row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            
            file_output = f"Laporan_SKP_{bulan_tahun.replace(' ', '_')}.docx"
            doc.save(file_output)
            
            with open(file_output, "rb") as fp:
                st.download_button(
                    label="📥 Unduh File Laporan Word (.docx)",
                    data=fp,
                    file_name=file_output,
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
        else:
            st.error("Belum ada data yang diinputkan!")
                
